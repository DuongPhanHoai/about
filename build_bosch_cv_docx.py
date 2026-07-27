#!/usr/bin/env python3
"""Build the Bosch-targeted CV using the uploaded DOCX as a visual template."""

from __future__ import annotations

import re
import sys
from pathlib import Path

from docx import Document
from docx.enum.style import WD_STYLE_TYPE
from docx.enum.text import WD_BREAK, WD_LINE_SPACING
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor


BLUE = RGBColor(31, 78, 121)
DARK = RGBColor(35, 35, 35)
GRAY = RGBColor(89, 89, 89)


def set_cell_border(cell, **edges):
    tc_pr = cell._tc.get_or_add_tcPr()
    borders = tc_pr.first_child_found_in("w:tcBorders")
    if borders is None:
        borders = OxmlElement("w:tcBorders")
        tc_pr.append(borders)
    for edge, attributes in edges.items():
        tag = f"w:{edge}"
        element = borders.find(qn(tag))
        if element is None:
            element = OxmlElement(tag)
            borders.append(element)
        for key, value in attributes.items():
            element.set(qn(f"w:{key}"), str(value))


def configure_styles(document: Document) -> None:
    styles = document.styles
    normal = styles["Normal"]
    normal.font.name = "Arial"
    normal.font.size = Pt(9)
    normal.font.color.rgb = DARK
    normal.paragraph_format.space_after = Pt(2)
    normal.paragraph_format.line_spacing_rule = WD_LINE_SPACING.SINGLE

    heading_specs = {
        "Heading 1": (12, 7, 3),
        "Heading 2": (10.5, 5, 2),
        "Heading 3": (9.5, 3, 1),
    }
    for name, (size, before, after) in heading_specs.items():
        style = styles[name]
        style.font.name = "Arial"
        style.font.size = Pt(size)
        style.font.bold = True
        style.font.color.rgb = BLUE
        style.paragraph_format.space_before = Pt(before)
        style.paragraph_format.space_after = Pt(after)
        style.paragraph_format.keep_with_next = True

    if "CV Bullet" not in styles:
        bullet_style = styles.add_style("CV Bullet", WD_STYLE_TYPE.PARAGRAPH)
    else:
        bullet_style = styles["CV Bullet"]
    bullet_style.base_style = normal
    bullet_style.font.name = "Arial"
    bullet_style.font.size = Pt(9)
    bullet_style.paragraph_format.left_indent = Inches(0.18)
    bullet_style.paragraph_format.first_line_indent = Inches(-0.12)
    bullet_style.paragraph_format.space_after = Pt(1.5)
    bullet_style.paragraph_format.keep_together = True


def clear_body_except_header_table(document: Document) -> None:
    body = document._element.body
    kept_table = False
    for child in list(body):
        if child.tag == qn("w:sectPr"):
            continue
        if child.tag == qn("w:tbl") and not kept_table:
            kept_table = True
            continue
        body.remove(child)


def format_header(document: Document) -> None:
    if not document.tables:
        table = document.add_table(rows=1, cols=2)
    else:
        table = document.tables[0]

    table.autofit = False
    table.columns[0].width = Inches(1.02)
    table.columns[1].width = Inches(6.25)
    for cell in table.rows[0].cells:
        set_cell_border(
            cell,
            top={"val": "nil"},
            bottom={"val": "nil"},
            left={"val": "nil"},
            right={"val": "nil"},
            insideH={"val": "nil"},
            insideV={"val": "nil"},
        )

    cell = table.cell(0, 1)
    cell.text = ""
    title = cell.paragraphs[0]
    title.paragraph_format.space_after = Pt(2)
    run = title.add_run(
        "Senior Technical & Engineering Leader | "
        "Software Platform Delivery | System Integration & Performance"
    )
    run.font.name = "Arial"
    run.font.size = Pt(12)
    run.font.bold = True
    run.font.color.rgb = BLUE

    contact = cell.add_paragraph()
    contact.paragraph_format.space_after = Pt(1)
    contact_run = contact.add_run(
        "duongphanhoai@gmail.com  |  +84 909 994 004  |  Vietnam"
    )
    contact_run.font.name = "Arial"
    contact_run.font.size = Pt(8.5)
    contact_run.font.color.rgb = DARK

    links = cell.add_paragraph()
    links.paragraph_format.space_after = Pt(0)
    link_run = links.add_run(
        "linkedin.com/in/duong-phan-hoai  |  github.com/DuongPhanHoai"
    )
    link_run.font.name = "Arial"
    link_run.font.size = Pt(8)
    link_run.font.color.rgb = BLUE


def add_inline_markdown(paragraph, text: str) -> None:
    """Render the small Markdown subset used by the CV."""
    tokens = re.split(r"(\*\*.*?\*\*)", text)
    for token in tokens:
        if not token:
            continue
        bold = token.startswith("**") and token.endswith("**")
        value = token[2:-2] if bold else token
        run = paragraph.add_run(value)
        run.font.name = "Arial"
        run.font.size = Pt(9)
        run.font.bold = bold
        run.font.color.rgb = DARK


def add_cv_content(document: Document, markdown_path: Path) -> None:
    lines = markdown_path.read_text(encoding="utf-8").splitlines()
    start = next(i for i, line in enumerate(lines) if line == "## Professional Summary")

    for raw_line in lines[start:]:
        line = raw_line.strip()
        if not line or line == "---":
            continue

        if line.startswith("## "):
            paragraph = document.add_paragraph(style="Heading 1")
            paragraph.add_run(line[3:])
        elif line.startswith("### "):
            paragraph = document.add_paragraph(style="Heading 2")
            paragraph.add_run(line[4:])
        elif line.startswith("#### "):
            paragraph = document.add_paragraph(style="Heading 3")
            paragraph.add_run(line[5:])
        elif line.startswith("- "):
            paragraph = document.add_paragraph(style="CV Bullet")
            bullet = paragraph.add_run("• ")
            bullet.font.name = "Arial"
            bullet.font.size = Pt(9)
            add_inline_markdown(paragraph, line[2:])
        else:
            paragraph = document.add_paragraph()
            add_inline_markdown(paragraph, line)
            if line.startswith("**") and "|" in line:
                paragraph.paragraph_format.keep_with_next = True


def configure_page(document: Document) -> None:
    section = document.sections[0]
    section.top_margin = Inches(0.45)
    section.bottom_margin = Inches(0.45)
    section.left_margin = Inches(0.55)
    section.right_margin = Inches(0.55)
    section.header_distance = Inches(0.2)
    section.footer_distance = Inches(0.2)

    footer = section.footer
    footer_paragraph = footer.paragraphs[0]
    footer_paragraph.alignment = 2
    run = footer_paragraph.add_run("Duong Phan Hoai")
    run.font.name = "Arial"
    run.font.size = Pt(7)
    run.font.color.rgb = GRAY


def remove_trailing_empty_paragraph(document: Document) -> None:
    if not document.paragraphs:
        return
    paragraph = document.paragraphs[-1]
    if not paragraph.text:
        paragraph._element.getparent().remove(paragraph._element)


def main() -> None:
    if len(sys.argv) != 4:
        raise SystemExit(
            "Usage: build_bosch_cv_docx.py SOURCE.docx CONTENT.md OUTPUT.docx"
        )

    source, markdown, output = map(Path, sys.argv[1:])
    document = Document(source)
    clear_body_except_header_table(document)
    configure_styles(document)
    configure_page(document)
    format_header(document)
    add_cv_content(document, markdown)
    remove_trailing_empty_paragraph(document)

    document.core_properties.title = (
        "Duong Phan Hoai - Senior Technical and Engineering Leader"
    )
    document.core_properties.subject = (
        "Application for Bosch Senior Technical Expert - Automotive HPC"
    )
    document.core_properties.author = "Duong Phan Hoai"
    document.save(output)


if __name__ == "__main__":
    main()
